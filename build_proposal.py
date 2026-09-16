# -*- coding: utf-8 -*-
"""
Builds the research proposal DOCX strictly following:
Tengeru Institute of Community Development (TICD), Department of Postgraduate
Studies, Research and Consultancy, "Research Proposal / Dissertation Writing
Guidelines", April 2019 - Part A (3.1 components) and Part C (formatting).

Format: A4; Times New Roman; 14pt headings / 12pt text; double line spacing;
left margin 4 cm; right, top and bottom margins 2.54 cm; header/footer 1.27 cm;
Arabic page numbers at the bottom centre of the body; title page unnumbered;
APA author-date citations; references listed alphabetically with all authors.
"""
import math
import re
import sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from proposal_content import BLOCKS, REFERENCES, TITLE

# ------------------------------------------------------------- build modes --
# python3 build_proposal.py                -> full proposal (Ch. 1-3 + budget
#                                             + action plan + 25 references)
# python3 build_proposal.py --chapter-one  -> standalone Chapter One document
#                                             (title page + Ch. 1 + only the
#                                             references cited in Chapter One)
CHAPTER_ONE_ONLY = "--chapter-one" in sys.argv

# Corporate authors are cited in the text by their abbreviation, so the matcher
# needs both the full name and the short form.
CORPORATE_TOKENS = {
    "United Republic of Tanzania": ["URT", "United Republic of Tanzania"],
    "National Bureau of Statistics": ["NBS", "National Bureau of Statistics"],
    "Office of the Controller and Auditor General":
        ["CAG", "Controller and Auditor General"],
    "Organisation for Economic Co-operation and Development":
        ["OECD", "Organisation for Economic Co-operation"],
    "Lushoto District Council": ["Lushoto District Council"],
}


def ref_tokens(ref):
    """Name tokens a reference can be cited by in the text."""
    head = ref.split("(")[0].strip().rstrip(".").strip()
    if head in CORPORATE_TOKENS:
        return CORPORATE_TOKENS[head]
    surnames = []
    for part in re.split(r",\s*&|,", head):
        part = part.strip()
        if part:
            surnames.append(part.split()[0].rstrip("."))
    return surnames[:1] or [head]


def ref_year(ref):
    m = re.search(r"\((\d{4})\)", ref) or re.search(r"(\d{4})", ref)
    return m.group(1)


def is_cited(ref, text):
    """True when the text cites this reference: an author name and the
    publication year must occur in the same citation window. Matching the two
    independently would misfire on financial years (2021/22) and on works
    cited elsewhere in the document."""
    year = ref_year(ref)
    ypat = (r"(?:\(|,\s|;\s|et al\.,\s|and\s)" + year + r"(?![\d/\-])"
            r"|" + year + r"\s*\)")
    for tok in ref_tokens(ref):
        for m in re.finditer(r"\b" + re.escape(tok) + r"\b", text):
            window = text[max(0, m.start() - 90): m.end() + 90]
            if re.search(ypat, window):
                return True
    return False


def chapter_one_blocks():
    """Title-page body blocks for a standalone Chapter One document: the
    Chapter One segment plus a reference list limited to works it cites."""
    start = next(i for i, (k, v) in enumerate(BLOCKS)
                 if k == "h1" and v.upper().startswith("CHAPTER ONE"))
    stop = next(i for i, (k, v) in enumerate(BLOCKS)
                if k == "h1" and v.upper().startswith("CHAPTER TWO"))
    segment = BLOCKS[start:stop]
    while segment and segment[0][0] == "pagebreak":
        segment = segment[1:]
    text = " ".join(v for k, v in segment if k in ("p", "h2", "h3", "caption", "source"))
    cited = [r for r in REFERENCES if is_cited(r, text)]
    return segment + [("h1", "REFERENCES")] + [("ref", r) for r in cited], cited


if CHAPTER_ONE_ONLY:
    RENDER_BLOCKS, CITED_REFS = chapter_one_blocks()
else:
    RENDER_BLOCKS, CITED_REFS = BLOCKS, REFERENCES

FONT = "Times New Roman"
PAGE_H_PT = (29.7 - 2.54 - 2.54) * 28.3465    # usable text height in points (top margin 1 inch)
PAGE_W_CM = 21.0 - 4.0 - 2.54                   # usable text width in cm


# ------------------------------------------------------------------ helpers --
def set_run(run, size=12, bold=False, italic=False, caps=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), FONT)
    if caps:
        c = OxmlElement('w:caps'); c.set(qn('w:val'), '1'); rpr.append(c)


def add_rich(par, text, size=12, bold=False, caps=False):
    """Adds text, honouring *italic* markers."""
    for i, chunk in enumerate(re.split(r'\*', text)):
        if not chunk:
            continue
        run = par.add_run(chunk)
        set_run(run, size=size, bold=bold, italic=(i % 2 == 1), caps=caps)


def fmt_par(par, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=2.0,
            before=0, after=0, indent=None, hanging=None, keep_next=False):
    pf = par.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    pf.keep_with_next = keep_next
    pf.widow_control = True


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0]
    par.text = ""
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    set_run(run, size=12)
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def restart_numbering(section, start=1):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), str(start))
    pgNumType.set(qn('w:fmt'), 'decimal')


def set_cell_borders(cell, on=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single' if on else 'nil')
        el.set(qn('w:sz'), '6' if on else '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tcPr.append(borders)


# ---------------------------------------------------------------- document ---
doc = Document()

# Base style
normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
normal.paragraph_format.line_spacing = 2.0
normal.paragraph_format.space_after = Pt(0)

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(4.0)
sec.right_margin = Cm(2.54)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.header_distance = Cm(1.27)
sec.footer_distance = Cm(1.27)
sec.different_first_page_header_footer = True

# ---- title page (section 1, no page number) ----
def title_page_par(text, size=14, bold=True, caps=False, before=0, after=0, spacing=1.5):
    par = doc.add_paragraph()
    fmt_par(par, size=size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=spacing,
            before=before, after=after)
    add_rich(par, text, size=size, bold=bold, caps=caps)
    return par

title_page_par("[NAME OF INSTITUTION]", size=14, before=0, after=6)
title_page_par("[FACULTY / DEPARTMENT]", size=14, after=6)
title_page_par("[PROGRAMME OF STUDY]", size=14, after=60)
title_page_par(TITLE, size=14, caps=False, after=60)
title_page_par("BY", size=14, after=12)
title_page_par("[STUDENT'S FULL NAME]", size=14, after=6)
title_page_par("[REGISTRATION NUMBER]", size=14, after=60)
p = doc.add_paragraph()
fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, after=48)
add_rich(p, "A RESEARCH PROPOSAL SUBMITTED IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE "
            "AWARD OF [DEGREE TITLE] OF [NAME OF INSTITUTION]", size=12, bold=False)
title_page_par("[MONTH], 2026", size=14, before=24)

# ---- new section for the body: Arabic page numbers starting at 1 ----
new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
new_sec.page_width = Cm(21.0); new_sec.page_height = Cm(29.7)
new_sec.left_margin = Cm(4.0); new_sec.right_margin = Cm(2.54)
new_sec.top_margin = Cm(2.54); new_sec.bottom_margin = Cm(2.54)
new_sec.header_distance = Cm(1.27); new_sec.footer_distance = Cm(1.27)
# doc.add_section() copies the title section's sectPr, which carries
# w:titlePg; clear it so Word does not suppress the number on body page 1
new_sec.different_first_page_header_footer = False
add_page_number(new_sec)
restart_numbering(new_sec, 1)


def add_heading1(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=4, after=2,
            keep_next=True)
    add_rich(par, text.upper(), size=14, bold=True)


def add_heading2(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, before=6, after=0,
            keep_next=True)
    add_rich(par, text, size=14, bold=True)


def add_heading3(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, before=4, after=0,
            keep_next=True)
    add_rich(par, text, size=12, bold=True)


def add_body(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=2.0)
    add_rich(par, text, size=12)


def add_caption(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, before=4, after=2,
            keep_next=True)
    add_rich(par, text, size=12, bold=True)


def add_source(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, before=2, after=4)
    add_rich(par, text, size=12)


def add_reference(text):
    par = doc.add_paragraph()
    fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=2.0, indent=1.27, hanging=1.27)
    add_rich(par, text, size=12)


def add_table(spec):
    rows = spec["rows"]
    widths = spec.get("widths")
    merge = spec.get("merge", False)
    header = spec.get("header", True)
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        is_head = header and ri == 0
        arrow_row = (not header) and all(c.strip() in ("\u2193", "") for c in row)
        for ci in range(ncols):
            cell = cells[ci]
            text = row[ci] if ci < len(row) else ""
            if widths and ci < len(widths):
                cell.width = Cm(widths[ci])
            par = cell.paragraphs[0]
            par.text = ""
            align = WD_ALIGN_PARAGRAPH.CENTER if (is_head or arrow_row or len(text) < 12) \
                else WD_ALIGN_PARAGRAPH.LEFT
            fmt_par(par, align=align, spacing=1.0, before=1, after=1)
            lines = text.split("\n")
            for li, line in enumerate(lines):
                target = par if li == 0 else cell.add_paragraph()
                if li > 0:
                    fmt_par(target, align=align, spacing=1.0, before=0, after=1)
                add_rich(target, line, size=12, bold=is_head)
        if merge and ri in (0, 2, 3, 4, 5) and len(row) >= 3:
            cells[0].merge(cells[-1])
    # keep table rows from splitting across pages
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trPr.append(cant)
    return tbl


# ------------------------------------------------------------ render blocks --
first_block = True
for kind, payload in RENDER_BLOCKS:
    if kind == "pagebreak":
        if first_block:
            first_block = False
            continue
        par = doc.add_paragraph()
        par.add_run().add_break(WD_BREAK.PAGE)
        continue
    first_block = False
    if kind == "h1":
        add_heading1(payload)
    elif kind == "h2":
        add_heading2(payload)
    elif kind == "h3":
        add_heading3(payload)
    elif kind == "p":
        add_body(payload)
    elif kind == "caption":
        add_caption(payload)
    elif kind == "source":
        add_source(payload)
    elif kind == "ref":
        add_reference(payload)
    elif kind == "table":
        add_table(payload)

if CHAPTER_ONE_ONLY:
    OUT = "Chapter-One-ME-System-and-Local-Government-Project-Performance-Lushoto.docx"
else:
    OUT = "Research-Proposal-ME-System-and-Local-Government-Project-Performance-Lushoto.docx"
doc.save(OUT)
print("saved:", OUT)


# ------------------------------------------------------------- page estimate --
# Single source of truth for pagination: estimate_pages.simulate(), which wraps
# text with real Times AFM glyph widths at the exact text width (14.46 cm) and
# models Word's widow/orphan control and keep-with-next behaviour.
import estimate_pages
from proposal_content import REFERENCES as _REFS

PAGE_H_PT = estimate_pages.PAGE_H
pages, per = estimate_pages.simulate(RENDER_BLOCKS, verbose=False)
print("\nEstimated layout (A4, top margin 1 inch -> usable text height %.0f pt)" % PAGE_H_PT)
for k, v in per.items():
    print("  %-52s %6.0f pt  %5.2f pages" % (k[:52], v, v / PAGE_H_PT))
print("  %-52s %6.0f pt  %5.2f pages" % ("BODY (packed)", sum(per.values()), sum(per.values()) / PAGE_H_PT))
print("  TOTAL = %d pages (1 title page + %d body pages)" % (pages, pages - 1))
if CHAPTER_ONE_ONLY:
    print("  mode: standalone CHAPTER ONE (title page + Ch. 1 + references cited in Ch. 1)")
    print("  references cited in Chapter One: %d of %d" % (len(CITED_REFS), len(_REFS)))
else:
    print("  guide limit: not less than 15 and not more than 20 pages ->",
          "OK" if 15 <= pages <= 20 else "OUT OF RANGE")
    print("  references: %d (guide requires 20-25) -> %s"
          % (len(_REFS), "OK" if 20 <= len(_REFS) <= 25 else "OUT OF RANGE"))
print("  body prose words: %d" % sum(len(v.replace("*", "").split()) for k, v in RENDER_BLOCKS if k == "p"))
